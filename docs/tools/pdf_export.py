# -*- coding: utf-8 -*-
"""PDF-копии книг: листы xlsx → HTML → печать в PDF через headless Chrome.

Excel и LibreOffice для этого не нужны. Значения берутся из сохранённого файла
(те же, что видит просмотрщик), оформление — из стилей ячеек.
"""
import base64
import html
import os
import shutil
import signal
import subprocess
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
PAGE_W = 1500.0     # ширина области печати A3 альбомной при полях 10 мм, px


def available():
    return os.path.exists(CHROME)


def fmt(v, nf):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if nf in ("0%", "0.0%"):
        n = 0 if nf == "0%" else 1
        return ("%.*f" % (n, v * 100)).replace(".", ",") + "%"
    if nf.startswith("+0.0"):
        return "0,0" if abs(v) < 0.05 else ("%+.1f" % v).replace(".", ",").replace("-", "−")
    if nf == "#,##0":
        return "{:,.0f}".format(v).replace(",", " ")
    if nf in ("0", "0.0", "0.00"):
        n = len(nf) - 2 if "." in nf else 0
        x = 0.0 if abs(v) < 0.5 * 10 ** (-n) else v
        return ("%.*f" % (n, x)).replace(".", ",").replace("-", "−")
    if isinstance(v, int) or float(v).is_integer():
        return str(int(v))
    return ("%.2f" % v).replace(".", ",")


def _rgb(color):
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and len(rgb) == 8 and rgb[2:] != "000000":
        return "#" + rgb[2:]
    return None


def _side(side):
    if side is None or side.style is None:
        return None
    if side.style == "medium":
        return "2px solid #000"
    if side.style == "hair":
        return "1px solid #bfbfbf"
    return "1px solid #000"


def _css(c):
    out = []
    if c.fill is not None and c.fill.fill_type == "solid":
        bg = _rgb(c.fill.fgColor)
        if bg:
            out.append("background:%s" % bg)
    f = c.font
    if f is not None:
        if f.b:
            out.append("font-weight:bold")
        if f.i:
            out.append("font-style:italic")
        col = _rgb(f.color) if f.color is not None else None
        if col:
            out.append("color:%s" % col)
        if f.sz and abs(f.sz - 11) > 0.1:
            out.append("font-size:%spt" % f.sz)
    a = c.alignment
    if a is not None:
        if a.horizontal in ("center", "right"):
            out.append("text-align:%s" % a.horizontal)
        if a.vertical == "top":
            out.append("vertical-align:top")
    b = c.border
    if b is not None:
        for name in ("left", "right", "top", "bottom"):
            s = _side(getattr(b, name))
            if s:
                out.append("border-%s:%s" % (name, s))
    return ";".join(out)


def _sheet_html(ws, image=None):
    n_rows, n_cols = ws.max_row, ws.max_column
    if image is not None:
        n_rows = min(n_rows, 2)
    widths = [(ws.column_dimensions[get_column_letter(c)].width or 9) * 7.2 + 6 for c in range(1, n_cols + 1)]
    total = sum(widths)
    merged, skip = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    skip.add((r, c))
    rows = []
    for r in range(1, n_rows + 1):
        cells, c = [], 1
        while c <= n_cols:
            if (r, c) in skip:
                c += 1
                continue
            cell = ws.cell(r, c)
            rs, cs = merged.get((r, c), (1, 1))
            text = fmt(cell.value, cell.number_format)
            if (r, c) not in merged and isinstance(cell.value, str) and len(text) > 12:
                # как в Excel: длинный текст занимает пустые ячейки справа
                k = c + 1
                while k <= n_cols and ws.cell(r, k).value is None and (r, k) not in skip and (r, k) not in merged \
                        and ws.cell(r, k).border.left.style is None and ws.cell(r, k).fill.fill_type is None:
                    k += 1
                cs = k - c
            attrs = ""
            if rs > 1:
                attrs += ' rowspan="%d"' % rs
            if cs > 1:
                attrs += ' colspan="%d"' % cs
            style = _css(cell)
            cells.append("<td%s%s>%s</td>" % (attrs, ' style="%s"' % style if style else "", html.escape(text).replace("\n", "<br>")))
            c += cs
        h = ws.row_dimensions[r].height
        rows.append("<tr%s>%s</tr>" % (' style="height:%dpx"' % round(h * 1.33) if h else "", "".join(cells)))
    cols = "".join('<col style="width:%.0fpx">' % w for w in widths)
    table = '<table style="width:%.0fpx"><colgroup>%s</colgroup>%s</table>' % (total, cols, "".join(rows))
    zoom = min(1.0, PAGE_W / total)
    body = '<div style="zoom:%.3f">%s</div>' % (zoom, table)
    if image is not None:
        data = base64.b64encode(Path(image).read_bytes()).decode("ascii")
        body += '<img src="data:image/png;base64,%s">' % data
    return '<section><div class="tab">%s</div>%s</section>' % (html.escape(ws.title), body)


_FONTS = "/Applications/Microsoft Word.app/Contents/Resources/DFonts"
_FACES = "".join(
    "@font-face { font-family: Calibri; src: url('file://%s/%s'); font-weight: %s; font-style: %s; }\n" % (
        _FONTS.replace(" ", "%20"), name, weight, style)
    for name, weight, style in (("Calibri.ttf", "normal", "normal"), ("Calibrib.ttf", "bold", "normal"),
                                ("Calibrii.ttf", "normal", "italic"))
    if os.path.exists(os.path.join(_FONTS, name)))

CSS = _FACES + """
@page { size: A3 landscape; margin: 10mm; }
* { box-sizing: border-box; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: Calibri, Arial, Helvetica, sans-serif; font-size: 11pt; color: #000; margin: 0; }
section { page-break-after: always; }
section:last-child { page-break-after: auto; }
.tab { font-size: 9pt; color: #7f7f7f; margin-bottom: 4px; }
table { border-collapse: collapse; table-layout: fixed; }
td { padding: 2px 4px; vertical-align: middle; overflow-wrap: break-word; line-height: 1.2; }
tr { page-break-inside: avoid; }
img { max-width: 100%; max-height: 250mm; display: block; margin-top: 6px; }
"""


def export(xlsx_path, pdf_path, images=None):
    """images: {имя листа: путь к png} — листы-рисунки."""
    images = images or {}
    wb = load_workbook(xlsx_path, data_only=True)
    parts = [_sheet_html(ws, images.get(ws.title)) for ws in wb.worksheets]
    doc = "<!doctype html><html lang='ru'><head><meta charset='utf-8'><title>%s</title><style>%s</style></head><body>%s</body></html>" % (
        html.escape(Path(xlsx_path).stem), CSS, "".join(parts))
    tmp = tempfile.mkdtemp(prefix="studbook_pdf_")
    try:
        page = Path(tmp) / "book.html"
        page.write_text(doc, encoding="utf-8")
        out = Path(tmp) / "book.pdf"
        proc = subprocess.Popen([CHROME, "--headless=new", "--disable-gpu", "--no-first-run", "--no-pdf-header-footer", "--allow-file-access-from-files",
                                 "--user-data-dir=%s" % (Path(tmp) / "profile"), "--print-to-pdf=%s" % out, page.as_uri()],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        # Chrome после печати может не завершиться сам: ждём готовый файл и закрываем процесс
        size, stable, deadline = -1, 0, time.time() + 90
        while time.time() < deadline and stable < 3:
            time.sleep(0.5)
            if proc.poll() is not None and out.exists():
                break
            now = out.stat().st_size if out.exists() else -1
            stable = stable + 1 if (now > 1000 and now == size) else 0
            size = now
        try:                                   # закрываем Chrome вместе со вспомогательными процессами
            os.killpg(proc.pid, signal.SIGTERM)
        except (ProcessLookupError, PermissionError):
            pass
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()
        if not out.exists() or out.stat().st_size < 1000:
            raise RuntimeError("Chrome не создал PDF")
        shutil.move(str(out), str(pdf_path))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return pdf_path
