# -*- coding: utf-8 -*-
"""Книга 2: «Расписание» (фаза 3, пункт 3)."""
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Border, Side

import inputs as I
import model as M
from book_pert import dmy, money, ru
from xlsx_util import *  # noqa: F401,F403

SH_TAB, SH_GANTT, SH_RES, SH_LOAD = "Расписание", "Гант", "Ресурсы", "Загрузка"
SH_CONF, SH_MS, SH_FIG1, SH_FIG2 = "Конфликты", "Вехи", "Диаграмма Ганта", "Загрузка (график)"
IDS = list(M.PK)
BAR = PatternFill("solid", fgColor="7F7F7F")
RED = PatternFill("solid", fgColor="C00000")
GREY = PatternFill("solid", fgColor="D9D9D9")
BUF = PatternFill("solid", fgColor="FFE699")


def who(codes):
    return ", ".join(I.PEOPLE[c]["name"].split()[0] for c in codes)


ORDER = ["Фаза 3", "Фаза 4", "В1", "В2", "В3", "В4", "после В4"]


def busy_text(plan, pid):
    """«Иванов, Петров заняты в 3.2, 3.5» — кто и в каких пакетах был занят."""
    groups = {}
    for other, xs in plan.waited_for(pid).items():
        groups.setdefault(tuple(xs), []).append(other)
    return "; ".join("%s заняты в %s" % (who(xs), ", ".join(pk)) for xs, pk in groups.items())


def ms_contents(plan):
    """Какие пакеты закрываются к каждой паре по расписанию."""
    out = {code: [] for code, _, _, _ in M.MILESTONES}
    for pid in plan.disc:
        code = plan.milestone(pid)
        if code in out:
            out[code].append(pid)
    return out


def sheet_table(wb, plan):
    ws = wb.create_sheet(SH_TAB)
    page(ws)
    title(ws, "Расписание после выравнивания ресурсов (вариант Б)",
          "Работы с 19.10.2026. Рабочие дни: пн–пт, кроме 04.11.2026.", 15)
    header(ws, 4, ["№", "Пакет работ", "Тип", "Исполнители", "E ост., ч", "D, дн", "Раннее начало по сети",
                   "Готовы предшественницы", "Старт", "Финиш", "Задержка из-за ресурсов, дн", "Причина задержки",
                   "Веха (план)", "Веха (устав)", "Крит."], height=48)
    for k, pid in enumerate(IDS):
        r, p = 5 + k, M.PK[pid]
        off, kind = pid in plan.cut, M.PK[pid]["kind"]
        chain = (not off) and kind == "disc" and pid in plan.res_chain
        muted = off or kind == "done"
        fill = FILL_LOE if (kind == "loe" and not off) else FILL_CRIT if chain else None
        cell(ws, "A%d" % r, pid, fill=fill, align=CENTER, font=F_MUTED if muted else F_BASE)
        cell(ws, "B%d" % r, p["name"], fill=fill, font=F_MUTED if muted else F_BASE)
        if muted:
            cell(ws, "C%d" % r, "исключён" if off else "выполнена", align=CENTER, font=F_MUTED)
            for col in "DEFGHIJKLM":
                cell(ws, "%s%d" % (col, r), "—", align=CENTER, font=F_MUTED)
            cell(ws, "N%d" % r, p["charter_ms"], align=CENTER, font=F_MUTED)
            cell(ws, "O%d" % r, "", align=CENTER)
            continue
        cell(ws, "D%d" % r, who(plan.team[pid]))
        cell(ws, "E%d" % r, M.E_rem(pid), fmt="0.0", align=CENTER)
        if kind == "loe":
            a, b = plan.win[pid]
            cell(ws, "C%d" % r, "гамак", align=CENTER)
            cell(ws, "F%d" % r, b - a, fmt="0.0", align=CENTER)
            cell(ws, "G%d" % r, "—", align=CENTER)
            cell(ws, "H%d" % r, "—", align=CENTER)
            cell(ws, "I%d" % r, M.start_date(a), align=CENTER)
            cell(ws, "J%d" % r, M.finish_date(a, b), align=CENTER)
            for col in "KLM":
                cell(ws, "%s%d" % (col, r), "—", align=CENTER)
            cell(ws, "N%d" % r, p["charter_ms"], align=CENTER)
            cell(ws, "O%d" % r, "", align=CENTER)
            continue
        s, f = plan.S[pid], plan.F[pid]
        ready = plan.block[pid][0]
        wait = s - ready if pid != "1.4" else 0.0
        cell(ws, "C%d" % r, "обычная", align=CENTER)
        cell(ws, "F%d" % r, plan.dur(pid), fmt="0.0", align=CENTER)
        cell(ws, "G%d" % r, M.start_date(plan.ES[pid]), align=CENTER)
        cell(ws, "H%d" % r, M.start_date(ready), align=CENTER)
        cell(ws, "I%d" % r, M.start_date(s), align=CENTER, font=F_BOLD)
        cell(ws, "J%d" % r, M.finish_date(s, f), align=CENTER, font=F_BOLD)
        cell(ws, "K%d" % r, wait if wait > 1e-6 else 0, fmt="0.0", align=CENTER)
        if pid == "1.4":
            cause = "после финальной пары 26.12"
        elif wait > 1e-6:
            cause = busy_text(plan, pid)
        else:
            cause = ""
        cell(ws, "L%d" % r, cause)
        code = plan.milestone(pid)
        late = p["charter_ms"] in ("В1", "В2", "В3") and ORDER.index(code) > ORDER.index(p["charter_ms"])
        cell(ws, "M%d" % r, code, align=CENTER, fill=FILL_WARN if late else None)
        cell(ws, "N%d" % r, p["charter_ms"], align=CENTER)
        cell(ws, "O%d" % r, "да" if chain else "", align=CENTER, fill=FILL_CRIT if chain else None)
    for r in range(5, 5 + len(IDS)):
        ws.row_dimensions[r].height = 30
    r = 6 + len(IDS)
    ws.cell(r, 2, "Критическая цепочка с учётом ресурсов: " + " → ".join(plan.res_chain)).font = F_BOLD
    ws.cell(r + 1, 2, "Жёлтым отмечены вехи позже, чем в уставе.").font = F_SMALL
    widths(ws, [6, 40, 12, 30, 8, 7, 12, 18, 12, 12, 13, 40, 10, 12, 7])
    ws.freeze_panes = "C5"


def sheet_gantt(wb, plan):
    ws = wb.create_sheet(SH_GANTT)
    page(ws)
    n = M.T_CLOSE
    title(ws, "Диаграмма Ганта (вариант Б), по рабочим дням",
          "Красный — критическая цепочка, тёмно-серый — остальные работы, светло-серый — гамаки, жёлтый — запас.", 5 + n)
    header(ws, 4, ["№", "Пакет работ", "Исполнители", "Старт", "Финиш"], height=18)
    for c in range(1, 6):
        ws.merge_cells(start_row=4, start_column=c, end_row=6, end_column=c)
    col0 = 6
    thick = Side(style="medium", color="000000")
    ms_cols = {t - 1: code for code, _, t, _ in M.MILESTONES if t > 0}
    for mon, i0, i1 in M.weeks(n):
        c = ws.cell(4, col0 + i0, M.CAL[i0].strftime("%d.%m"))
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER
        ws.merge_cells(start_row=4, start_column=col0 + i0, end_row=4, end_column=col0 + i1 - 1)
    for i in range(n):
        col = get_column_letter(col0 + i)
        ws.column_dimensions[col].width = 2.9
        c = ws.cell(5, col0 + i, M.CAL[i].day)
        c.font, c.alignment = Font(name=FONT, size=8), CENTER
        m = ws.cell(6, col0 + i, ms_cols.get(i, ""))
        m.font, m.alignment = Font(name=FONT, size=8, bold=True), CENTER
    ws.row_dimensions[6].height = 14
    chain = set(plan.res_chain)
    for k, pid in enumerate(IDS):
        r, p = 7 + k, M.PK[pid]
        off, kind = pid in plan.cut, p["kind"]
        muted = off or kind == "done"
        cell(ws, "A%d" % r, pid, align=CENTER, font=F_MUTED if muted else F_BASE)
        cell(ws, "B%d" % r, p["name"], font=F_MUTED if muted else F_BASE)
        if muted:
            cell(ws, "C%d" % r, "исключён" if off else "выполнена", font=F_MUTED)
            cell(ws, "D%d" % r, "—", align=CENTER, font=F_MUTED)
            cell(ws, "E%d" % r, "—", align=CENTER, font=F_MUTED)
            a = b = None
        else:
            a, b = plan.win[pid] if kind == "loe" else (plan.S[pid], plan.F[pid])
            cell(ws, "C%d" % r, who(plan.team[pid]), font=F_SMALL)
            cell(ws, "D%d" % r, M.start_date(a), fmt="DD.MM", align=CENTER)
            cell(ws, "E%d" % r, M.finish_date(a, b), fmt="DD.MM", align=CENTER)
        for i in range(n):
            c = ws.cell(r, col0 + i)
            if a is not None and M.rnd(a) <= i <= max(M.rnd(b) - 1, M.rnd(a)):
                c.fill = GREY if kind == "loe" else RED if pid in chain else BAR
            elif pid == "6.3" and a is not None and M.rnd(b) <= i < M.T_ACC:
                c.fill = BUF
            c.border = Border(right=thick) if i in ms_cols else Border(right=Side(style="hair", color="BFBFBF")) if M.CAL[i].weekday() == 4 else Border()
        ws.row_dimensions[r].height = 16
    r = 8 + len(IDS)
    ws.cell(r, 2, "Жирная линия — последний рабочий день перед парой: " +
            ", ".join("%s %s" % (code, dt.strftime("%d.%m")) for code, dt, t, _ in M.MILESTONES if t > 0) + ".").font = F_SMALL
    widths(ws, [6, 40, 28, 8, 8])
    ws.freeze_panes = "F7"


def sheet_resources(wb, reg, plan):
    ws = wb.create_sheet(SH_RES)
    page(ws)
    title(ws, "Ресурсы (вариант Б)", None, 12)
    header(ws, 3, ["Код", "Участник", "Роль", "Ч/нед по уставу", "Ч/нед в плане", "Из них на гамаки", "Ставка, USD/ч",
                   "Часов по плану с 19.10", "Фонд 19.10–31.12, ч", "Загрузка", "Стоимость с 19.10, USD"], height=48)
    for k, x in enumerate(I.ALL):
        r, rec = 4 + k, I.PEOPLE[x]
        res = plan.reserve.get(x, 0) if x != "QA" else "всё время с начала 3.1"
        hrs = plan.person_hours(x)
        vals = [x, rec["name"], rec["role"], rec["h"], plan.hours[x], res, rec["rate"], hrs,
                M.T_CLOSE * plan.hours[x] / 5.0, hrs / (M.T_CLOSE * plan.hours[x] / 5.0), hrs * rec["rate"]]
        fmts = [None, None, None, "0", "0", "0.0", "0", "0.0", "0.0", "0%", "#,##0"]
        for c, (v, f) in enumerate(zip(vals, fmts), 1):
            cell(ws, "%s%d" % (get_column_letter(c), r), v, fmt=f if not isinstance(v, str) else None,
                 align=WRAP if c in (2, 3, 6) else CENTER)
    r = 4 + len(I.ALL)
    cell(ws, "A%d" % r, "")
    cell(ws, "B%d" % r, "Итого", font=F_BOLD)
    for col in "CFG":
        cell(ws, "%s%d" % (col, r), "")
    for col, exp, fmt in (("D", 54, "0"), ("E", plan.week_h, "0"), ("H", plan.rem_h, "0.0"), ("I", plan.fund_rem, "0.0"), ("K", plan.cost_rem, "#,##0")):
        reg.formula(ws, "%s%d" % (col, r), "=SUM(%s4:%s%d)" % (col, col, r - 1), expect=exp, fmt=fmt, font=F_BOLD, align=CENTER)
    reg.formula(ws, "J%d" % r, "=H%d/I%d" % (r, r), expect=plan.rem_h / plan.fund_rem, fmt="0%", font=F_BOLD, align=CENTER)

    r0 = r + 3
    ws.cell(r0 - 1, 1, "Часы по пакетам и исполнителям (остаток работ с 19.10.2026)").font = F_H2
    header(ws, r0, ["№", "Пакет работ", "Тип"] + list(I.ALL) + ["Итого ост., ч", "Выполнено до 17.10, ч", "Стоимость пакета, USD"], height=48)
    scope = [p for p in IDS if p not in plan.cut]
    for k, pid in enumerate(scope):
        rr, kind = r0 + 1 + k, M.PK[pid]["kind"]
        cell(ws, "A%d" % rr, pid, align=CENTER)
        cell(ws, "B%d" % rr, M.PK[pid]["name"])
        cell(ws, "C%d" % rr, {"done": "выполнена", "loe": "гамак", "disc": "обычная"}[kind], align=CENTER)
        for j, x in enumerate(I.ALL):
            v = plan.H[pid].get(x, 0)
            cell(ws, "%s%d" % (get_column_letter(4 + j), rr), v if v > 1e-9 else None, fmt="0.0", align=CENTER)
        reg.formula(ws, "K%d" % rr, "=SUM(D%d:J%d)" % (rr, rr), expect=M.E_rem(pid), fmt="0.0", align=CENTER)
        cell(ws, "L%d" % rr, M.E_done(pid) if M.E_done(pid) > 0 else None, fmt="0.0", align=CENTER)
        cell(ws, "M%d" % rr, plan.pkg_cost(pid), fmt="#,##0", align=CENTER)
    rt = r0 + 1 + len(scope)
    cell(ws, "A%d" % rt, "")
    cell(ws, "B%d" % rt, "Итого", font=F_BOLD)
    cell(ws, "C%d" % rt, "")
    for j, x in enumerate(I.ALL):
        col = get_column_letter(4 + j)
        reg.formula(ws, "%s%d" % (col, rt), "=SUM(%s%d:%s%d)" % (col, r0 + 1, col, rt - 1), expect=plan.person_hours(x),
                    fmt="0.0", font=F_BOLD, align=CENTER)
    for col, exp, fmt in (("K", plan.rem_h, "0.0"), ("L", plan.done_h, "0.0"), ("M", plan.bac, "#,##0")):
        reg.formula(ws, "%s%d" % (col, rt), "=SUM(%s%d:%s%d)" % (col, r0 + 1, col, rt - 1), expect=exp, fmt=fmt,
                    font=F_BOLD, align=CENTER)
    ws.cell(rt + 2, 1, "Стоимость пакета = часы исполнителей × ставки, включая выполненную часть. С резервом 10%%: %s USD." % money(round(plan.bac) * 1.1)).font = F_SMALL
    widths(ws, [6, 40, 24, 9, 9, 12, 9, 10, 11, 10, 12, 12, 12])
    ws.freeze_panes = "A4"


def sheet_load(wb, plan):
    ws = wb.create_sheet(SH_LOAD)
    page(ws)
    wk, after, cap = plan.load(True)
    wk0, before, cap0 = plan.load(False)
    n = len(wk)
    title(ws, "Загрузка по неделям, % от недельного фонда",
          "«До» — работы начинаются в раннее начало по сети, «после» — по расписанию.", 3 + n)
    header(ws, 4, ["Участник", "Ч/нед", ""] + [w[0].strftime("%d.%m") for w in wk] + ["Макс."], height=22)
    r = 5
    for x in I.ALL:
        cell(ws, "A%d" % r, I.PEOPLE[x]["name"])
        cell(ws, "B%d" % r, plan.hours[x], fmt="0", align=CENTER)
        cell(ws, "A%d" % (r + 1), None)
        cell(ws, "B%d" % (r + 1), None)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)
        for row, label, data, caps in ((r, "до", before, cap0), (r + 1, "после", after, cap)):
            cell(ws, "C%d" % row, label, align=CENTER)
            vals = [(data[x][k] / caps[x][k]) if k < len(data[x]) else 0 for k in range(n)]
            for k, v in enumerate(vals):
                cell(ws, "%s%d" % (get_column_letter(4 + k), row), v, fmt="0%", align=CENTER, fill=FILL_BAD if v > 1.005 else None)
            cell(ws, "%s%d" % (get_column_letter(4 + n), row), max(vals), fmt="0%", align=CENTER, font=F_BOLD,
                 fill=FILL_BAD if max(vals) > 1.005 else None)
        r += 2
    cell(ws, "A%d" % r, "Команда, ч", font=F_BOLD)
    cell(ws, "B%d" % r, plan.week_h, fmt="0", align=CENTER)
    cell(ws, "C%d" % r, "после", align=CENTER)
    for k in range(n):
        cell(ws, "%s%d" % (get_column_letter(4 + k), r), sum(after[x][k] for x in I.ALL), fmt="0", align=CENTER)
    cell(ws, "%s%d" % (get_column_letter(4 + n), r), max(sum(after[x][k] for x in I.ALL) for k in range(n)), fmt="0", align=CENTER, font=F_BOLD)
    over = sum(1 for x in I.ALL for k in range(min(n, len(before[x]))) if before[x][k] / cap0[x][k] > 1.005)
    ws.cell(r + 2, 1, "Перегруз (больше 100%%): до выравнивания %d человеко-недель, после — 0." % over).font = F_BASE
    widths(ws, [18, 7, 7] + [8] * n + [8])
    ws.freeze_panes = "D5"
    return over


def sheet_conflicts(wb, plan):
    ws = wb.create_sheet(SH_CONF)
    page(ws)
    title(ws, "Ресурсные конфликты и их разрешение (вариант Б)", None, 7)
    header(ws, 3, ["Пакет", "Название", "Мог начаться", "Начинается", "Задержка, раб. дн", "Причина", "Полный резерв по сети, дн"], height=34)
    rows = [(pid, plan.S[pid] - plan.block[pid][0]) for pid in plan.disc if pid != "1.4" and plan.S[pid] - plan.block[pid][0] > 1e-6]
    rows.sort(key=lambda t: plan.S[t[0]])
    for k, (pid, wait) in enumerate(rows):
        r = 4 + k
        cell(ws, "A%d" % r, pid, align=CENTER)
        cell(ws, "B%d" % r, M.PK[pid]["name"])
        cell(ws, "C%d" % r, M.start_date(plan.block[pid][0]), align=CENTER)
        cell(ws, "D%d" % r, M.start_date(plan.S[pid]), align=CENTER)
        cell(ws, "E%d" % r, wait, fmt="0.0", align=CENTER)
        cell(ws, "F%d" % r, busy_text(plan, pid))
        cell(ws, "G%d" % r, plan.TF[pid], fmt="0.0", align=CENTER)
        ws.row_dimensions[r].height = 30
    r = 5 + len(rows)
    ws.cell(r, 1, "Правило: один человек одновременно в одном пакете, раньше идёт работа с меньшим поздним началом.").font = F_BASE
    ws.cell(r + 1, 1, "Окончание приёмки 6.3: по сети %s раб. дн, по расписанию %s раб. дн. Запас до 25.12 — %s раб. дн." % (
        ru(plan.T_net), ru(plan.T_sched), ru(plan.buffer))).font = F_BASE
    widths(ws, [8, 40, 14, 14, 12, 50, 14])
    ws.freeze_panes = "A4"
    return len(rows)


def sheet_milestones(wb, reg, plan):
    ws = wb.create_sheet(SH_MS)
    page(ws)
    title(ws, "Вехи и плановый объём (PV) нарастающим итогом", None, 8)
    header(ws, 3, ["Веха", "Дата пары", "По уставу (п. 3.1)", "Готово по расписанию", "PV, ч", "PV, USD", "Доля", "Прирост, USD"], height=22)
    cont = ms_contents(plan)
    prev = None
    rows = [(code, dt, t, text) for code, dt, t, text in M.MILESTONES] + [("Закрытие", I.CLOSE_BY, M.T_CLOSE, "Закрытие проекта")]
    for k, (code, dt, t, text) in enumerate(rows):
        r = 4 + k
        h, usd = plan.pv(t)
        if code == "Фаза 3":
            ready = "1.1; 1.2 на 60%; 2.1 на 70%"
        elif code == "Закрытие":
            ready = "1.4"
        else:
            ready = "; ".join("%s %s" % (p, M.PK[p]["short"]) for p in sorted(cont.get(code, [])))
        cell(ws, "A%d" % r, code, align=CENTER)
        cell(ws, "B%d" % r, dt, align=CENTER)
        cell(ws, "C%d" % r, text)
        cell(ws, "D%d" % r, ready)
        cell(ws, "E%d" % r, h, fmt="0", align=CENTER)
        cell(ws, "F%d" % r, usd, fmt="#,##0", align=CENTER)
        reg.formula(ws, "G%d" % r, "=F%d/F%d" % (r, 3 + len(rows)), expect=usd / plan.bac, fmt="0%", align=CENTER)
        if prev is None:
            cell(ws, "H%d" % r, usd, fmt="#,##0", align=CENTER)
        else:
            reg.formula(ws, "H%d" % r, "=F%d-F%d" % (r, r - 1), expect=usd - prev, fmt="#,##0", align=CENTER)
        prev = usd
        ws.row_dimensions[r].height = 16 * max(2, len(ready) // 42 + 1, len(text) // 40 + 1)
    r = 5 + len(rows)
    ws.cell(r, 1, "BAC = %s USD, с резервом 10%% — %s USD. PV считается на конец последнего рабочего дня перед парой." % (
        money(plan.bac), money(round(plan.bac) * 1.1))).font = F_BASE
    widths(ws, [10, 12, 46, 52, 8, 11, 8, 12])
    ws.freeze_panes = "A4"


def build(path, plan, gantt_png, load_png):
    wb = Workbook()
    wb.remove(wb.active)
    reg = Registry()
    sheet_table(wb, plan)
    sheet_gantt(wb, plan)
    sheet_resources(wb, reg, plan)
    sheet_load(wb, plan)
    sheet_conflicts(wb, plan)
    sheet_milestones(wb, reg, plan)
    for name, png in ((SH_FIG1, gantt_png), (SH_FIG2, load_png)):
        ws = wb.create_sheet(name)
        page(ws)
        img = XLImage(str(png))
        scale = 1500.0 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, "A1")
    wb.calculation.fullCalcOnLoad = True
    values = verify(wb, reg)
    wb.save(path)
    n = inject_cached(path, wb, values)
    return len(reg.items), n
