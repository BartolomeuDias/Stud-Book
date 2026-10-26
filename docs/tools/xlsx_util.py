# -*- coding: utf-8 -*-
"""Оформление листов, проверка формул и запись их значений в файл.

openpyxl не вычисляет формулы, поэтому после сохранения в файле нет значений —
просмотрщики без пересчёта (Quick Look, превью в облаке) показали бы пустые ячейки.
Здесь формулы вычисляются собственным мини-интерпретатором, сверяются с моделью
и записываются в xlsx как кэшированные значения.
"""
import datetime as _dt
import math
import re
import shutil
import zipfile
from xml.sax.saxutils import escape

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

FONT = "Calibri"
INK, INK2, MUTED = "000000", "000000", "7F7F7F"
F_BASE = Font(name=FONT, size=11, color=INK)
F_BOLD = Font(name=FONT, size=11, bold=True, color=INK)
F_SMALL = Font(name=FONT, size=10, color=INK)
F_NOTE = Font(name=FONT, size=11, color=INK)
F_HEAD = Font(name=FONT, size=11, bold=True, color=INK)
F_TITLE = Font(name=FONT, size=14, bold=True, color=INK)
F_H2 = Font(name=FONT, size=11, bold=True, color=INK)
F_INPUT = Font(name=FONT, size=11, color=INK)
F_MUTED = Font(name=FONT, size=11, color=MUTED)

FILL_HEAD = PatternFill("solid", fgColor="D9D9D9")
FILL_INPUT = PatternFill()
FILL_CRIT = PatternFill("solid", fgColor="FFC7CE")
FILL_LOE = PatternFill("solid", fgColor="EDEDED")
FILL_OFF = PatternFill()
FILL_NOTE = PatternFill()
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
FILL_KEY = PatternFill()

_S = Side(style="thin", color="000000")
BORDER = Border(left=_S, right=_S, top=_S, bottom=_S)
WRAP = Alignment(wrap_text=True, vertical="center")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
RIGHT = Alignment(vertical="center", horizontal="right")


class Registry:
    """Запоминает формулы и ожидаемые моделью значения."""

    def __init__(self):
        self.items = []     # (лист, ячейка, ожидаемое значение или None)

    def formula(self, ws, ref, text, expect=None, **style):
        ws[ref] = text
        self.items.append((ws.title, ref, expect))
        cell(ws, ref, **style)
        return ws[ref]


def cell(ws, ref, value=None, font=None, fill=None, fmt=None, align=None, border=True):
    if isinstance(value, (_dt.date, _dt.datetime)):
        # даты пишутся текстом: формат «ДД.ММ.ГГГГ» часть просмотрщиков показывает как номер дня в году
        value, fmt = value.strftime("%d.%m" if fmt == "DD.MM" else "%d.%m.%Y"), None
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font or F_BASE
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align or WRAP
    if border:
        c.border = BORDER
    return c


def header(ws, row, titles, col0=1, height=32):
    for k, t in enumerate(titles):
        cell(ws, "%s%d" % (get_column_letter(col0 + k), row), t, font=F_HEAD, fill=FILL_HEAD, align=CENTER)
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None, width=10):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 22
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws.row_dimensions[2].height = 18


def text_block(ws, row, text, width, height=None, fill=None, font=None, bold=False):
    c = ws.cell(row, 1, text)
    c.font = font or (F_BOLD if bold else F_BASE)
    c.alignment = WRAP_TOP
    if fill is not None:
        c.fill = fill
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    if height is None:
        chars = sum(ws.column_dimensions[get_column_letter(k)].width or 10 for k in range(1, width + 1))
        height = 15 * max(1, int(math.ceil(len(text) * 1.15 / max(chars, 1))))
    ws.row_dimensions[row].height = height
    return c


def widths(ws, values):
    for k, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(k)].width = w


def page(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.zoomScale = 100


# --- вычисление формул -------------------------------------------------------
def _nums(args):
    out = []
    for a in args:
        if isinstance(a, (list, tuple)):
            out.extend(_nums(a))
        elif isinstance(a, bool):
            continue
        elif isinstance(a, (int, float)):
            out.append(a)
    return out


FUNCS = {
    "MAX": lambda *a: max(_nums(a)) if _nums(a) else 0,
    "MIN": lambda *a: min(_nums(a)) if _nums(a) else 0,
    "SUM": lambda *a: sum(_nums(a)),
    "ABS": abs,
    "SQRT": math.sqrt,
    "ROUND": lambda x, n=0: round(x, int(n)) + 0.0,
    "IF": lambda c, a, b=0: a if c else b,
    "NORMSDIST": lambda z: 0.5 * (1.0 + math.erf(z / math.sqrt(2.0))),
}

_TOKEN = re.compile(
    r"""(?P<str>"[^"]*")
      |(?P<xref>'[^']+'!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
      |(?P<func>[A-Z]+)(?=\()
      |(?P<ref>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
      |(?P<num>\d+(?:\.\d+)?)
      |(?P<op><>|<=|>=|=|<|>|\^|[-+*/(),])
      |(?P<ws>\s+)""", re.X)


class Calc:
    def __init__(self, wb):
        self.wb, self.memo = wb, {}

    def get(self, sheet, ref):
        key = (sheet, ref.replace("$", ""))
        if key not in self.memo:
            v = self.wb[sheet][key[1]].value
            if isinstance(v, str) and v.startswith("="):
                v = self.evaluate(sheet, v[1:])
            self.memo[key] = 0 if v is None else v
        return self.memo[key]

    def rng(self, sheet, a, b):
        c0, r0, c1, r1 = range_boundaries("%s:%s" % (a.replace("$", ""), b.replace("$", "")))
        return [self.get(sheet, "%s%d" % (get_column_letter(c), r)) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]

    def _ref(self, sheet, r):
        if ":" in r:
            a, b = r.split(":")
            return "_r(%r,%r,%r)" % (sheet, a, b)
        return "_c(%r,%r)" % (sheet, r)

    def evaluate(self, sheet, expr):
        out, pos = [], 0
        while pos < len(expr):
            m = _TOKEN.match(expr, pos)
            if not m:
                raise ValueError("не разобрана формула: %s (позиция %d)" % (expr, pos))
            kind, tok, pos = m.lastgroup, m.group(), m.end()
            if kind == "str":
                out.append(repr(tok[1:-1]))
            elif kind == "xref":
                sh, r = tok.rsplit("!", 1)
                out.append(self._ref(sh[1:-1], r))
            elif kind == "ref":
                out.append(self._ref(sheet, tok))
            elif kind == "func":
                out.append("_F[%r]" % tok)
            elif kind == "num":
                out.append(tok)
            elif kind == "op":
                out.append({"=": "==", "<>": "!=", "^": "**"}.get(tok, tok))
        return eval("".join(out), {"_F": FUNCS, "_c": self.get, "_r": self.rng})


def verify(wb, reg, tol=1e-6):
    """Вычисляет все зарегистрированные формулы и сверяет с ожидаемыми значениями модели."""
    calc, values, errors = Calc(wb), {}, []
    for sheet, ref, expect in reg.items:
        v = calc.get(sheet, ref)
        values.setdefault(sheet, {})[ref] = v
        if expect is not None:
            ok = (v == expect) if isinstance(expect, str) else (isinstance(v, (int, float)) and abs(v - expect) <= tol * max(1.0, abs(expect)))
            if not ok:
                errors.append("%s!%s: формула даёт %r, модель %r" % (sheet, ref, v, expect))
    if errors:
        raise AssertionError("Формулы расходятся с моделью:\n  " + "\n  ".join(errors[:30]))
    return values


_CELL = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)>(<f>.*?</f>)<v\s*/>|<c r="([A-Z]+\d+)"([^>]*)>(<f>.*?</f>)<v></v>', re.S)


def inject_cached(path, wb, values):
    """Дописывает в сохранённый xlsx вычисленные значения формул."""
    tmp = str(path) + ".tmp"
    names = {"xl/worksheets/sheet%d.xml" % (i + 1): ws.title for i, ws in enumerate(wb.worksheets)}
    done = 0
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in names:
                vals = values.get(names[item.filename], {})
                counter = [0]

                def repl(m):
                    ref, attrs, f = (m.group(1), m.group(2), m.group(3)) if m.group(1) else (m.group(4), m.group(5), m.group(6))
                    if ref not in vals:
                        return m.group(0)
                    v = vals[ref]
                    counter[0] += 1
                    if isinstance(v, str):
                        return '<c r="%s"%s t="str">%s<v>%s</v>' % (ref, attrs, f, escape(v))
                    if isinstance(v, bool):
                        return '<c r="%s"%s t="b">%s<v>%d</v>' % (ref, attrs, f, int(v))
                    return '<c r="%s"%s>%s<v>%s</v>' % (ref, attrs, f, repr(float(v) + 0.0 if v else 0.0))

                data = _CELL.sub(repl, data.decode("utf-8")).encode("utf-8")
                done += counter[0]
            zout.writestr(item, data)
    shutil.move(tmp, str(path))
    return done
