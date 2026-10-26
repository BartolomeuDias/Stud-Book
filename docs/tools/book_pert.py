# -*- coding: utf-8 -*-
"""Книга 1: «Сетевой график и PERT» (фаза 3, пункт 2)."""
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

import inputs as I
import model as M
from xlsx_util import *  # noqa: F401,F403

SH_DATA, SH_A, SH_B = "Данные", "PERT (вариант А)", "PERT (вариант Б)"
SH_EV, SH_PROB, SH_SC, SH_DEC, SH_ASM, SH_FIG = "События", "Вероятность", "Сценарии", "Изменения", "Допущения", "Сетевой график"

IDS = list(M.PK)
ROW = {pid: 5 + k for k, pid in enumerate(IDS)}     # одинаковые строки на «Данные» и листах PERT
TOTAL = 5 + len(IDS)                                # 30
SUMROW = TOTAL + 1                                  # 31 — Tкр


def ru(x, n=1):
    x = 0.0 if abs(x) < 0.5 * 10 ** (-n) else x
    return ("%.*f" % (n, x)).replace(".", ",")


def pct(x, n=1):
    return ru(100 * x, n) + "%"


def money(x):
    return "{:,.0f}".format(x).replace(",", " ")


def dmy(d):
    return d.strftime("%d.%m.%Y")


def kind_ru(plan, pid):
    if pid in plan.cut:
        return "исключён"
    return {"done": "выполнена", "loe": "гамак", "disc": "обычная"}[M.PK[pid]["kind"]]


def net_preds_text(pid, plan_b):
    p = M.PK[pid]
    if p["kind"] == "done":
        return "—"
    if p["kind"] == "loe":
        return "гамак"
    if pid == "6.3":
        return ", ".join(plan_b.preds["6.3"])
    return ", ".join(p["preds"]) if p["preds"] else "—"


def sheet_data(wb, reg, plan_b):
    ws = wb.create_sheet(SH_DATA)
    page(ws)
    title(ws, "Исходные данные по пакетам работ",
          "Статус-дата 17.10.2026. O, M, P взяты из сводной таблицы оценок.", 17)
    header(ws, 4, ["№", "Пакет работ", "Приоритет", "O, ч", "M, ч", "P, ч", "E, ч", "σ, ч", "Готово на 17.10",
                   "E ост., ч", "σ ост., ч", "Тип", "Предшественницы (сводная таблица)",
                   "Предшественницы (в сети)", "Примечание", "В плане Б", "E ост. в плане Б, ч"], height=46)
    for pid in IDS:
        p, r = M.PK[pid], ROW[pid]
        off = pid in plan_b.cut
        fill = FILL_LOE if p["kind"] == "loe" else None
        font = F_MUTED if off else F_BASE
        cell(ws, "A%d" % r, pid, align=CENTER, fill=fill, font=font)
        cell(ws, "B%d" % r, p["name"], fill=fill, font=font)
        cell(ws, "C%d" % r, p["moscow"], align=CENTER)
        for col, key in (("D", "o"), ("E", "m"), ("F", "p")):
            cell(ws, "%s%d" % (col, r), p[key], fmt="0", align=CENTER)
        reg.formula(ws, "G%d" % r, "=(D%d+4*E%d+F%d)/6" % (r, r, r), expect=M.E(pid), fmt="0.00", align=CENTER)
        reg.formula(ws, "H%d" % r, "=(F%d-D%d)/6" % (r, r), expect=M.sigma_h(pid), fmt="0.00", align=CENTER)
        cell(ws, "I%d" % r, p["done"], fmt="0%", align=CENTER)
        reg.formula(ws, "J%d" % r, "=G%d*(1-I%d)" % (r, r), expect=M.E_rem(pid), fmt="0.00", align=CENTER)
        reg.formula(ws, "K%d" % r, "=H%d*(1-I%d)" % (r, r), expect=M.sigma_rem(pid), fmt="0.00", align=CENTER)
        cell(ws, "L%d" % r, {"done": "выполнена", "loe": "гамак", "disc": "обычная"}[p["kind"]], align=CENTER)
        cell(ws, "M%d" % r, p["preds_table"], align=CENTER)
        cell(ws, "N%d" % r, net_preds_text(pid, plan_b), align=CENTER)
        cell(ws, "O%d" % r, I.PRED_NOTES.get(pid, ""))
        cell(ws, "P%d" % r, "нет" if off else "да", align=CENTER)
        reg.formula(ws, "Q%d" % r, '=IF(P%d="да",J%d,0)' % (r, r), expect=0.0 if off else M.E_rem(pid), fmt="0.00", align=CENTER)
        ws.row_dimensions[r].height = 30
    cell(ws, "A%d" % TOTAL, "")
    cell(ws, "B%d" % TOTAL, "Итого", font=F_BOLD)
    for col in "DEFGJQ":
        exp = {"G": sum(M.E(x) for x in IDS), "J": sum(M.E_rem(x) for x in IDS), "Q": plan_b.rem_h}.get(col)
        reg.formula(ws, "%s%d" % (col, TOTAL), "=SUM(%s5:%s%d)" % (col, col, TOTAL - 1), expect=exp,
                    font=F_BOLD, fmt="0.00" if col in "GJQ" else "0", align=CENTER)
    widths(ws, [6, 40, 11, 7, 7, 7, 8, 8, 10, 9, 9, 12, 22, 22, 38, 9, 12])
    ws.freeze_panes = "C5"


def sheet_pert(wb, reg, plan, name, head):
    ws = wb.create_sheet(name)
    page(ws)
    title(ws, head, "Сроки в рабочих днях от 19.10.2026.", 17)
    header(ws, 4, ["№", "Пакет работ", "Тип", "E ост., ч", "σ ост., ч", "Исполнители", "Мощность, ч/дн", "D, дн",
                   "σ, дн", "ES", "EF", "LS", "LF", "Rполн", "Rсвоб", "Rнез", "Крит."], height=34)
    D = "'%s'" % SH_DATA
    for pid in IDS:
        p, r = M.PK[pid], ROW[pid]
        off, kind = pid in plan.cut, M.PK[pid]["kind"]
        crit = (not off) and kind == "disc" and pid in plan.critical
        fill = FILL_LOE if (kind == "loe" and not off) else FILL_CRIT if crit else None
        muted = off or kind == "done"
        cell(ws, "A%d" % r, pid, fill=fill, align=CENTER, font=F_MUTED if muted else F_BASE)
        cell(ws, "B%d" % r, p["name"], fill=fill, font=F_MUTED if muted else F_BASE)
        cell(ws, "C%d" % r, kind_ru(plan, pid), align=CENTER, font=F_MUTED if muted else F_BASE)
        reg.formula(ws, "D%d" % r, "=%s!J%d" % (D, r), expect=M.E_rem(pid), fmt="0.00", align=CENTER)
        reg.formula(ws, "E%d" % r, "=%s!K%d" % (D, r), expect=M.sigma_rem(pid), fmt="0.00", align=CENTER)
        if muted:
            for col in "FGHIJKLMNOPQ":
                cell(ws, "%s%d" % (col, r), "—", align=CENTER, font=F_MUTED)
            continue
        cell(ws, "F%d" % r, ", ".join(plan.team[pid]), align=CENTER)
        if kind == "loe":
            spec = I.LOE[pid]
            a = "=0" if spec["start"] == 0 else "=J%d" % ROW[spec["start"].split(":")[1]]
            b = "=J%d" % ROW["1.4" if spec["end"] == "ACC" else spec["end"].split(":")[1]]
            w = plan.win_net[pid]
            cell(ws, "G%d" % r, "—", align=CENTER)
            reg.formula(ws, "J%d" % r, a, expect=w[0], fmt="0.00", align=CENTER)
            reg.formula(ws, "K%d" % r, b, expect=w[1], fmt="0.00", align=CENTER)
            reg.formula(ws, "H%d" % r, "=K%d-J%d" % (r, r), expect=w[1] - w[0], fmt="0.00", align=CENTER)
            for col in "ILMNOPQ":
                cell(ws, "%s%d" % (col, r), "—", align=CENTER)
            continue
        pr, su = plan.preds[pid], plan.succ[pid]
        cell(ws, "G%d" % r, plan.power(pid), fmt="0.00", align=CENTER)
        reg.formula(ws, "H%d" % r, "=D%d/G%d" % (r, r), expect=plan.dur(pid), fmt="0.00", align=CENTER)
        reg.formula(ws, "I%d" % r, "=E%d/G%d" % (r, r), expect=plan.sig_d(pid), fmt="0.00", align=CENTER)
        es = "=MAX(%s)" % ",".join("K%d" % ROW[q] for q in pr) if pr else "=0"
        reg.formula(ws, "J%d" % r, es, expect=plan.ES[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "K%d" % r, "=J%d+H%d" % (r, r), expect=plan.EF[pid], fmt="0.00", align=CENTER)
        lf = "=MIN(%s)" % ",".join("L%d" % ROW[q] for q in su) if su else "=K%d" % r
        reg.formula(ws, "M%d" % r, lf, expect=plan.LF[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "L%d" % r, "=ROUND(M%d-H%d,9)" % (r, r), expect=plan.LS[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "N%d" % r, "=ROUND(L%d-J%d,6)" % (r, r), expect=plan.TF[pid], fmt="0.00", align=CENTER,
                    fill=FILL_CRIT if crit else None)
        nxt = "MIN(%s)" % ",".join("J%d" % ROW[q] for q in su) if su else "M%d" % r
        prv = "MAX(%s)" % ",".join("M%d" % ROW[q] for q in pr) if pr else "0"
        reg.formula(ws, "O%d" % r, "=ROUND(%s-K%d,6)" % (nxt, r), expect=plan.FF[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "P%d" % r, "=ROUND(MAX(0,%s-%s-H%d),6)" % (nxt, prv, r), expect=plan.IF[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "Q%d" % r, '=IF(ABS(N%d)<0.000001,"да","")' % r, expect="да" if crit else "",
                    align=CENTER, fill=FILL_CRIT if crit else None)
    for r in range(5, TOTAL):
        ws.row_dimensions[r].height = 21
    r = SUMROW
    for k, (label, f, exp) in enumerate([
        ("Tкр (окончание финальной приёмки 6.3), раб. дн", "=K%d" % ROW["6.3"], plan.T_net),
        ("Окончание проекта (1.4), раб. дн", "=K%d" % ROW["1.4"], plan.T),
        ("Tдир (25.12.2026), раб. дн", None, M.T_ACC),
    ]):
        c = ws.cell(r + k, 2, label)
        c.font = F_BOLD
        if f:
            reg.formula(ws, "D%d" % (r + k), f, expect=exp, fmt="0.00", font=F_BOLD, align=CENTER)
        else:
            cell(ws, "D%d" % (r + k), exp, fmt="0", font=F_BOLD, align=CENTER)
    ws.cell(r + 3, 2, "Критический путь: " + " → ".join(plan.crit_path)).font = F_BOLD
    notes = [
        "D = E ост. / мощность; σ (дн) = σ ост. (ч) / мощность. ES = max EF предшественниц; EF = ES + D; LF = min LS последователей; LS = LF − D.",
        "Rполн = LS − ES; Rсвоб = min ES последователей − EF; Rнез = max(0; min ES последователей − max LF предшественниц − D).",
        "Гамак — растянутая работа: длительность берётся из сети, резервы не считаются.",
    ]
    for k, t in enumerate(notes):
        ws.cell(r + 5 + k, 2, t).font = F_SMALL
    widths(ws, [6, 42, 12, 10, 10, 20, 11, 8, 8, 8, 8, 8, 8, 8, 8, 8, 7])
    ws.freeze_panes = "C5"


def sheet_events(wb, reg, plan):
    ws = wb.create_sheet(SH_EV)
    page(ws)
    title(ws, "Ранние и поздние сроки событий (вариант Б)",
          "Событие — окончание работы: TE = EF, TL = LF. Рабочие дни от 19.10.2026.", 7)
    header(ws, 4, ["Событие", "Содержание", "Завершается работа", "Могут начаться", "TE", "TL", "Резерв"], height=32)
    cell(ws, "A5", 0, align=CENTER)
    cell(ws, "B5", "Начало оставшихся работ (статус-дата)")
    cell(ws, "C5", "—", align=CENTER)
    cell(ws, "D5", ", ".join(p for p in plan.disc if not plan.preds[p]), align=CENTER)
    cell(ws, "E5", 0, fmt="0.00", align=CENTER)
    cell(ws, "F5", 0, fmt="0.00", align=CENTER)
    reg.formula(ws, "G5", "=F5-E5", expect=0.0, fmt="0.00", align=CENTER)
    S = "'%s'" % SH_B
    order = sorted(plan.disc, key=lambda p: (plan.EF[p], p))
    for k, pid in enumerate(order, 1):
        r, src = 5 + k, ROW[pid]
        fill = FILL_CRIT if pid in plan.critical else None
        cell(ws, "A%d" % r, k, align=CENTER, fill=fill)
        name = "Проект закрыт" if pid == "1.4" else "Финальная приёмка пройдена" if pid == "6.3" else "Готово: " + M.PK[pid]["name"]
        cell(ws, "B%d" % r, name, fill=fill)
        cell(ws, "C%d" % r, pid, align=CENTER)
        cell(ws, "D%d" % r, ", ".join(plan.succ[pid]) or "—", align=CENTER)
        reg.formula(ws, "E%d" % r, "=%s!K%d" % (S, src), expect=plan.EF[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "F%d" % r, "=%s!M%d" % (S, src), expect=plan.LF[pid], fmt="0.00", align=CENTER)
        reg.formula(ws, "G%d" % r, "=ROUND(F%d-E%d,6)" % (r, r), expect=plan.LF[pid] - plan.EF[pid], fmt="0.00", align=CENTER, fill=fill)
        ws.row_dimensions[r].height = 20
    widths(ws, [10, 52, 14, 26, 9, 9, 9])
    ws.freeze_panes = "A5"


def sheet_prob(wb, reg, pa, pb):
    ws = wb.create_sheet(SH_PROB)
    page(ws)
    title(ws, "Вероятность выполнения в срок и сверка с уставом", None, 4)
    widths(ws, [50, 40, 40, 32])
    A, B = "'%s'" % SH_A, "'%s'" % SH_B

    def head(r, t):
        ws.cell(r, 1, t).font = F_H2
        header(ws, r + 1, ["Показатель", "Вариант А", "Вариант Б", "Формула"], height=20)

    def line(r, label, va, vb, how="", fmt="0.00", key=False, expect=(None, None)):
        cell(ws, "A%d" % r, label)
        for col, v, exp in (("B", va, expect[0]), ("C", vb, expect[1])):
            ref = "%s%d" % (col, r)
            if isinstance(v, str) and v.startswith("="):
                reg.formula(ws, ref, v, expect=exp, fmt=fmt, align=CENTER, font=F_BOLD if key else F_BASE)
            else:
                cell(ws, ref, v, fmt=fmt, align=CENTER, font=F_BOLD if key else F_BASE)
        cell(ws, "D%d" % r, how)
        ws.row_dimensions[r].height = 18

    def sig(sheet, plan):
        return "=SQRT(%s)" % "+".join("%s!I%d^2" % (sheet, ROW[p]) for p in plan.crit_to_acc)

    head(3, "Вероятность срока (по сети)")
    line(5, "Часов в неделю", pa.week_h, pb.week_h, fmt="0")
    line(6, "Пакетов в расчёте", len(pa.scope) - 1, len(pb.scope) - 1, fmt="0")
    line(7, "Tкр, раб. дн", "=%s!D%d" % (A, SUMROW), "=%s!D%d" % (B, SUMROW), "EF работы 6.3", key=True, expect=(pa.T_net, pb.T_net))
    line(8, "Критический путь", " → ".join(pa.crit_to_acc), " → ".join(pb.crit_to_acc), fmt="@")
    line(9, "σкр, раб. дн", sig(A, pa), sig(B, pb), "√Σσ² по критическому пути", expect=(pa.sigma_net, pb.sigma_net))
    line(10, "Tдир, раб. дн", M.T_ACC, M.T_ACC, "19.10–25.12.2026", fmt="0")
    line(11, "Tдир − Tкр, раб. дн", "=B10-B7", "=C10-C7", expect=(M.T_ACC - pa.T_net, M.T_ACC - pb.T_net))
    line(12, "z", "=(B10-B7)/B9", "=(C10-C7)/C9", "(Tдир − Tкр) / σкр",
         expect=((M.T_ACC - pa.T_net) / pa.sigma_net, (M.T_ACC - pb.T_net) / pb.sigma_net))
    line(13, "P(T ≤ Tдир)", "=NORMSDIST(B12)", "=NORMSDIST(C12)", "Ф(z)", fmt="0.0%", key=True, expect=(pa.p_net, pb.p_net))

    head(15, "Трудоёмкость и фонд времени")
    D = "'%s'" % SH_DATA
    tot = sum(M.E(x) for x in IDS)
    line(17, "ΣE всех пакетов, ч", "=%s!G%d" % (D, TOTAL), "=%s!G%d" % (D, TOTAL), expect=(tot, tot))
    line(18, "Фонд времени по уставу, ч", I.FUND_CHARTER_H, I.FUND_CHARTER_H, "54 ч/нед × 15 нед", fmt="0")
    line(19, "Запас фонда, ч", "=B18-B17", "=C18-C17", expect=(I.FUND_CHARTER_H - tot, I.FUND_CHARTER_H - tot))
    line(20, "Выполнено на 17.10, ч", "=%s!G%d-%s!J%d" % (D, TOTAL, D, TOTAL), "=%s!G%d-%s!J%d" % (D, TOTAL, D, TOTAL),
         expect=(pa.done_h, pb.done_h))
    line(21, "Остаток работ, ч", "=%s!J%d" % (D, TOTAL), "=%s!Q%d" % (D, TOTAL), key=True, expect=(pa.rem_h, pb.rem_h))
    line(22, "Рабочих дней 19.10–31.12", M.T_CLOSE, M.T_CLOSE, fmt="0")
    line(23, "Фонд оставшегося времени, ч", "=B22*B5/5", "=C22*C5/5", "дни × ч/нед / 5", key=True, expect=(pa.fund_rem, pb.fund_rem))
    line(24, "Запас (+) / дефицит (−), ч", "=B23-B21", "=C23-C21", expect=(pa.fund_rem - pa.rem_h, pb.fund_rem - pb.rem_h))
    line(25, "Загрузка фонда", "=B21/B23", "=C21/C23", fmt="0%", expect=(pa.rem_h / pa.fund_rem, pb.rem_h / pb.fund_rem))
    line(26, "Всего часов по плану", "=B20+B21", "=C20+C21", expect=(pa.done_h + pa.rem_h, pb.done_h + pb.rem_h))

    head(28, "Стоимость, USD")
    line(30, "Работы по ставкам исполнителей", round(pa.bac), round(pb.bac), "часы × ставка", fmt="#,##0", key=True)
    line(31, "Резерв на риски 10%", "=B30*0.1", "=C30*0.1", fmt="#,##0", expect=(round(pa.bac) * 0.1, round(pb.bac) * 0.1))
    line(32, "Итого с резервом", "=B30+B31", "=C30+C31", fmt="#,##0", key=True, expect=(round(pa.bac) * 1.1, round(pb.bac) * 1.1))
    line(33, "Потолок бюджета", I.CEILING_USD, I.CEILING_USD, fmt="#,##0")
    line(34, "Запас до потолка", "=B33-B32", "=C33-C32", fmt="#,##0",
         expect=(I.CEILING_USD - round(pa.bac) * 1.1, I.CEILING_USD - round(pb.bac) * 1.1))
    line(35, "По средней ставке 55 USD/ч", "=B26*%d" % I.AVG_RATE, "=C26*%d" % I.AVG_RATE, "как в сводной таблице", fmt="#,##0",
         expect=((pa.done_h + pa.rem_h) * I.AVG_RATE, (pb.done_h + pb.rem_h) * I.AVG_RATE))

    head(37, "После выравнивания ресурсов (расписание)")
    line(39, "Окончание 6.3", dmy(M.finish_date(pa.S["6.3"], pa.F["6.3"])), dmy(M.finish_date(pb.S["6.3"], pb.F["6.3"])), fmt="@", key=True)
    line(40, "Запас до 25.12, раб. дн", pa.buffer, pb.buffer, fmt="0.0")
    line(41, "P(успеть)", pa.p_sched, pb.p_sched, fmt="0.0%", key=True)

    ws.cell(43, 1, "Вывод").font = F_H2
    concl = [
        "Вариант А (все пакеты, 54 ч/нед) в срок не укладывается: Tкр = %s раб. дн при Tдир = %d, P = %s. "
        "Остаток работ %s ч больше фонда оставшегося времени %s ч." % (
            ru(pa.T_net), M.T_ACC, pct(pa.p_net), ru(pa.rem_h, 0), ru(pa.fund_rem, 0)),
        "Вариант Б (Must и 5.1, каждому +2 ч/нед): Tкр = %s раб. дн, P = %s. После выравнивания ресурсов "
        "приёмка заканчивается %s, запас %s раб. дн, P = %s." % (
            ru(pb.T_net), pct(pb.p_net), dmy(M.finish_date(pb.S["6.3"], pb.F["6.3"])), ru(pb.buffer), pct(pb.p_sched, 0)),
        "Стоимость варианта Б %s USD, с резервом %s USD. Изменения к уставу приведены на листе «Изменения»." % (
            money(pb.bac), money(round(pb.bac) * 1.1)),
    ]
    for k, t in enumerate(concl):
        text_block(ws, 44 + k, t, 4, height=32)
    ws.freeze_panes = "B3"


def sheet_scenarios(wb, results):
    ws = wb.create_sheet(SH_SC)
    page(ws)
    title(ws, "Сценарии: недельная нагрузка и состав пакетов", None, 12)
    header(ws, 3, ["Вариант", "Ч/нед", "Состав", "Остаток работ, ч", "Фонд 19.10–31.12, ч", "Загрузка фонда",
                   "Tкр, дн", "P по сети", "Окончание 6.3", "Запас до 25.12, дн", "P по расписанию", "Итог"], height=46)
    for k, res in enumerate(results):
        r, p = 4 + k, res["plan"]
        base = res["label"] == "Б"
        verdict = "укладывается" if p.buffer >= 3 else "на грани" if p.buffer >= 0 else "не укладывается"
        fill = FILL_OK if verdict == "укладывается" else FILL_WARN if verdict == "на грани" else FILL_BAD
        font = F_BOLD if base else F_BASE
        vals = [res["label"], p.week_h, res["name"], p.rem_h, p.fund_rem, p.rem_h / p.fund_rem, p.T_net, p.p_net,
                dmy(M.finish_date(p.S["6.3"], p.F["6.3"])), p.buffer, p.p_sched, verdict]
        fmts = [None, "0", None, "0", "0", "0%", "0.0", "0.0%", None, "+0.0;-0.0;0.0", "0.0%", None]
        for c, (v, f) in enumerate(zip(vals, fmts), 1):
            cell(ws, "%s%d" % (get_column_letter(c), r), v, fmt=f, font=font,
                 align=WRAP if c == 3 else CENTER, fill=fill if c == 12 else None)
        ws.row_dimensions[r].height = 20
    r = 5 + len(results)
    ws.cell(r, 1, "Запас от 3 раб. дн — укладывается, от 0 до 3 — на грани.").font = F_SMALL
    ws.cell(r + 1, 1, "При 54 ч/нед не помещаются даже Must-пакеты. При +2 ч/нед каждому помещаются Must и один пакет Should (5.1 или 4.3).").font = F_BASE
    widths(ws, [9, 8, 30, 11, 12, 11, 9, 10, 15, 12, 12, 17])
    ws.freeze_panes = "A4"


def sheet_changes(wb, pa, pb, results):
    ws = wb.create_sheet(SH_DEC)
    page(ws)
    title(ws, "Изменения к уставу (запрос на изменение № 1)", "Требуют согласования по п. 4.2 устава.", 6)
    header(ws, 4, ["№", "Что меняется", "По уставу", "Предлагается", "Согласует", "Последствия"], height=22)
    must54 = next(x["plan"] for x in results if x["plan"].extra == 0 and x["name"] == "Только Must")
    ms = {c: [] for c, _, _, _ in M.MILESTONES}
    for pid in pb.disc:
        code = pb.milestone(pid)
        if code in ms and pid not in ("6.3", "1.2", "2.2", "2.4"):
            ms[code].append(pid)
    rows = [
        ("Часы в неделю", "54 (10 / 8 / 10 / 8 / 6 / 6 / 6)",
         "68: каждому +2 ч с 19.10 по 31.12",
         "команда, куратор, заказчик",
         "Всего по плану %s ч (фонд 810 ч). Стоимость %s USD, с резервом %s USD." % (
             ru(pb.done_h + pb.rem_h, 0), money(pb.bac), money(round(pb.bac) * 1.1))),
        ("Пакет 5.4 (Could)", "по возможности", "исключить", "урезается первым по п. 3.3", "Снимается критерий приёмки 16."),
        ("Пакеты Should", "4.3, 5.1, 5.2, 5.3", "оставить 5.1; исключить 4.3, 5.2, 5.3", "заказчик",
         "Ц5 не выполняется, Ц4 частично. Сквозной сценарий: заявка, проверка совместимости, реестр. "
         "Критерии 9 и 10 снимаются, 8 и 11 уточняются."),
        ("Какой пакет Should оставить", "приоритеты А-14: очередь, заявки, счета",
         "5.1 (вторая роль пользователя — клиент). Замена с тем же сроком: 4.3", "заказчик",
         "С 5.1 нет очереди и календаря, с 4.3 нет роли клиента."),
        ("Содержание вех", "В1: реестр, размещение; В2: справки, ввод записей; В3: совместимость, заявки, очередь; В4: счета, приёмка",
         "; ".join("%s: %s" % (c, ", ".join(sorted(v))) for c, v in ms.items() if v and c.startswith("В")) + ", приёмка",
         "заказчик, куратор",
         "Справки показываются на В3 вместо В2, размещение на В2 вместо В1. Пакеты 4.2 и 5.1 показываются на финальной паре."),
        ("Если +2 ч/нед не согласованы", "—",
         "при 54 ч/нед Must-пакеты заканчиваются %s" % dmy(M.finish_date(must54.S["6.3"], must54.F["6.3"])), "заказчик, куратор",
         "Придётся урезать Must-пакеты или понизить уровень реализации."),
    ]
    for k, row in enumerate(rows):
        r = 5 + k
        cell(ws, "A%d" % r, k + 1, align=CENTER)
        for c, v in enumerate(row, 2):
            cell(ws, "%s%d" % (get_column_letter(c), r), v, align=WRAP_TOP)
        ws.row_dimensions[r].height = 16 * max(2, max(len(row[1]) // 30, len(row[2]) // 36, len(row[4]) // 48) + 1)
    widths(ws, [5, 26, 34, 40, 24, 54])
    ws.freeze_panes = "A5"


def sheet_assumptions(wb, pa, pb):
    ws = wb.create_sheet(SH_ASM)
    page(ws, landscape=False)
    title(ws, "Допущения расчёта", None, 2)
    widths(ws, [5, 125])
    items = [
        "Статус-дата 17.10.2026. Рассчитываются оставшиеся работы, отсчёт с 19.10.2026. Рабочие дни: пн–пт, кроме 04.11.2026.",
        "Готовность на статус-дату: 1.1 — 100%, 1.2 — 60%, 2.1 — 70%, 1.3 — 38% (по времени), остальные пакеты не начаты.",
        "O, M, P взяты из сводной таблицы оценок. E = (O + 4M + P) / 6, σ = (P − O) / 6. Для начатых пакетов остаток пропорционален готовности.",
        "D = E ост. / мощность. Мощность — сумма часов в день исполнителей пакета (ч/нед : 5) без часов, занятых гамаками.",
        "Гамаки (растянутые работы): 1.3, остаток 2.1, 6.2. Длительность берётся из сети. Под них отведено: руководитель проекта 4 ч/нед, "
        "аналитик 1 ч/нед, техред %s ч/нед (в варианте А %s ч/нед); тестировщик с начала 3.1 занят только тестированием." % (
            ru(pb.reserve["TW"], 0), ru(pa.reserve["TW"], 0)),
        "6.3 в сети — подготовка и проведение финальной приёмки. Tдир = %d раб. дн (25.12.2026, последний рабочий день перед парой 26.12). "
        "Закрытие 1.4 — 28–31.12.2026." % M.T_ACC,
        "Связи «окончание–начало» взяты из сводной таблицы оценок, отличия указаны на листе «Данные».",
        "Вероятность срока: z = (Tдир − Tкр) / σкр, σкр = √Σσ² по критическому пути, P = Ф(z).",
        "Расписание: один человек одновременно занят в одном пакете; при конфликте раньше идёт работа с меньшим поздним началом (файл «Расписание»).",
        "Стоимость — часы исполнителя × его ставка (устав, п. 5.5).",
        "В команде 7 человек (устав, п. 3.3).",
        "Сеть рассчитана формулами на листах PERT. Выравнивание ресурсов и рисунки выполнены программой на Python.",
    ]
    for k, text in enumerate(items):
        r = 3 + k
        cell(ws, "A%d" % r, k + 1, align=CENTER, border=False)
        cell(ws, "B%d" % r, text, align=WRAP_TOP, border=False)
        ws.row_dimensions[r].height = 16 * (len(text) // 118 + 1) + 2


def build(path, pa, pb, results, fig_path):
    wb = Workbook()
    wb.remove(wb.active)
    reg = Registry()
    sheet_data(wb, reg, pb)
    sheet_pert(wb, reg, pa, SH_A, "Расчёт сети. Вариант А: все пакеты, 54 ч/нед")
    sheet_pert(wb, reg, pb, SH_B, "Расчёт сети. Вариант Б: базовый план, 68 ч/нед")
    sheet_events(wb, reg, pb)
    sheet_prob(wb, reg, pa, pb)
    sheet_scenarios(wb, results)
    sheet_changes(wb, pa, pb, results)
    sheet_assumptions(wb, pa, pb)
    ws = wb.create_sheet(SH_FIG)
    page(ws)
    img = XLImage(str(fig_path))
    scale = 1500.0 / img.width
    img.width, img.height = int(img.width * scale), int(img.height * scale)
    ws.add_image(img, "A1")
    wb.calculation.fullCalcOnLoad = True
    values = verify(wb, reg)
    wb.save(path)
    n = inject_cached(path, wb, values)
    return len(reg.items), n
